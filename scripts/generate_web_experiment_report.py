from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\。\Desktop\webwork")
OUT = ROOT / "高级Web技术实验报告2405010511李鹏-优化版.docx"
ASSET_DIR = ROOT / "report_assets"


def font(name="SimSun"):
    try:
        return ImageFont.truetype("C:/Windows/Fonts/simsun.ttc", 22)
    except Exception:
        return ImageFont.load_default()


def set_run_font(run, size=12, bold=False, name="SimSun", color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph(paragraph, first_line=True, line=1.5, space_after=6):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line
    fmt.space_after = Pt(space_after)
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    for run in paragraph.runs:
        set_run_font(run)


def add_text(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.add_run(text)
    set_paragraph(p, first_line=first_line)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 1:
        p.style = "Heading 1"
        set_run_font(run, size=16, bold=True, name="SimHei")
    elif level == 2:
        p.style = "Heading 2"
        set_run_font(run, size=14, bold=True, name="SimHei")
    else:
        p.style = "Heading 3"
        set_run_font(run, size=12, bold=True, name="SimHei")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        set_paragraph(p, first_line=False, line=1.35, space_after=3)


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F4F6F8")
    cell._tc.get_or_add_tcPr().append(shading)
    p = cell.paragraphs[0]
    for line in code.strip("\n").splitlines():
        r = p.add_run(line)
        set_run_font(r, size=9, name="Consolas")
        p.add_run("\n")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=10)
    doc.add_paragraph()


def draw_diagram(path, title, nodes, edges):
    ASSET_DIR.mkdir(exist_ok=True)
    w, h = 1300, 720
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    f_title = font()
    f_node = font()
    d.rectangle((0, 0, w, 70), fill=(29, 78, 216))
    d.text((40, 20), title, fill="white", font=f_title)
    boxes = {}
    for key, x, y, bw, bh, fill in nodes:
        boxes[key] = (x, y, x + bw, y + bh)
        d.rounded_rectangle(boxes[key], radius=18, fill=fill, outline=(38, 38, 38), width=2)
        lines = key.split("\\n")
        for idx, line in enumerate(lines):
            d.text((x + 22, y + 22 + idx * 30), line, fill=(20, 20, 20), font=f_node)
    for a, b, label in edges:
        ax1, ay1, ax2, ay2 = boxes[a]
        bx1, by1, bx2, by2 = boxes[b]
        start = (ax2, (ay1 + ay2) // 2)
        end = (bx1, (by1 + by2) // 2)
        d.line((start, end), fill=(31, 41, 55), width=4)
        d.polygon([(end[0], end[1]), (end[0] - 14, end[1] - 9), (end[0] - 14, end[1] + 9)], fill=(31, 41, 55))
        if label:
            d.text(((start[0] + end[0]) // 2 - 80, start[1] - 32), label, fill=(31, 41, 55), font=f_node)
    img.save(path)


def build_assets():
    draw_diagram(
        ASSET_DIR / "experiment1_structure.png",
        "实验一：本地静态社交分享平台结构图",
        [
            ("index.html\\n页面结构", 60, 150, 230, 110, (219, 234, 254)),
            ("styles.css\\n布局与主题", 380, 150, 230, 110, (220, 252, 231)),
            ("app.js\\n交互逻辑", 700, 150, 230, 110, (254, 243, 199)),
            ("localStorage\\n本地数据", 1010, 150, 230, 110, (253, 230, 138)),
            ("动态列表\\n发布/评论/点赞", 380, 420, 270, 120, (224, 231, 255)),
            ("身份切换\\n游客/用户/管理员", 720, 420, 270, 120, (254, 226, 226)),
        ],
        [
            ("index.html\\n页面结构", "styles.css\\n布局与主题", "加载样式"),
            ("styles.css\\n布局与主题", "app.js\\n交互逻辑", "绑定事件"),
            ("app.js\\n交互逻辑", "localStorage\\n本地数据", "读写"),
            ("app.js\\n交互逻辑", "动态列表\\n发布/评论/点赞", "渲染"),
            ("app.js\\n交互逻辑", "身份切换\\n游客/用户/管理员", "权限判断"),
        ],
    )
    draw_diagram(
        ASSET_DIR / "experiment2_flow.png",
        "实验二：ShareFlow 前后端分离数据流",
        [
            ("Vue 页面\\nHome/Publish/Profile", 50, 160, 250, 120, (219, 234, 254)),
            ("Pinia + Router\\n状态与守卫", 360, 160, 250, 120, (220, 252, 231)),
            ("Axios API\\nAuthorization", 670, 160, 250, 120, (254, 243, 199)),
            ("Spring Controller\\nREST 接口", 980, 160, 250, 120, (224, 231, 255)),
            ("Service\\n业务规则", 360, 440, 250, 120, (254, 226, 226)),
            ("MyBatis Mapper\\nSQL 映射", 670, 440, 250, 120, (237, 233, 254)),
            ("MySQL\\n用户/帖子/评论等", 980, 440, 250, 120, (253, 230, 138)),
        ],
        [
            ("Vue 页面\\nHome/Publish/Profile", "Pinia + Router\\n状态与守卫", "操作"),
            ("Pinia + Router\\n状态与守卫", "Axios API\\nAuthorization", "请求"),
            ("Axios API\\nAuthorization", "Spring Controller\\nREST 接口", "/api"),
            ("Spring Controller\\nREST 接口", "Service\\n业务规则", "调用"),
            ("Service\\n业务规则", "MyBatis Mapper\\nSQL 映射", "持久化"),
            ("MyBatis Mapper\\nSQL 映射", "MySQL\\n用户/帖子/评论等", "读写"),
        ],
    )
    # AI evidence panel requested by the course note in prior report workflow.
    img = Image.new("RGB", (1300, 520), (248, 250, 252))
    d = ImageDraw.Draw(img)
    f = font()
    d.rounded_rectangle((40, 40, 1260, 480), radius=24, fill="white", outline=(148, 163, 184), width=2)
    lines = [
        "AI 辅助提示词记录（节选）",
        "请根据 C:\\Users\\。\\Desktop\\web 与 C:\\Users\\。\\Desktop\\webwork 的真实源码，",
        "结合《高级Web实验及实验报告》要求，生成一份包含两个实验的课程实验报告。",
        "报告需要每章包含任务目的、实验要求、实验过程与结果、心得体会，",
        "并补充项目结构图、数据流说明图、关键接口和源码依据。",
    ]
    y = 80
    for i, line in enumerate(lines):
        d.text((80, y), line, fill=(15, 23, 42), font=f)
        y += 62 if i == 0 else 54
    img.save(ASSET_DIR / "ai_prompt_evidence.png")


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.4)
    styles = doc.styles
    styles["Normal"].font.name = "SimSun"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    styles["Normal"].font.size = Pt(12)


def cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("湖南科技大学计算机科学与工程学院")
    set_run_font(r, size=18, bold=True, name="SimHei")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("高级WEB技术 实验报告")
    set_run_font(r, size=22, bold=True, name="SimHei")
    for _ in range(6):
        doc.add_paragraph()
    for label, value in [
        ("专业班级：", "计科五班"),
        ("姓    名：", "李鹏"),
        ("学    号：", "2405010511"),
        ("指导教师：", "陈向"),
        ("时    间：", "3.30-5.29"),
        ("地    点：", "逸夫楼222"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}          {value}")
        set_run_font(r, size=14, name="SimSun")
    doc.add_page_break()


def toc(doc):
    add_heading(doc, "目录", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("高级 Web 技术实验报告内容提纲")
    set_run_font(r, size=14, bold=True, name="SimHei")
    p.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(2.3)
    table.columns[1].width = Cm(10.2)
    table.columns[2].width = Cm(2.6)
    headers = ["序号", "章节内容", "页码"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EEF8")
        cell._tc.get_or_add_tcPr().append(shading)
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, size=11, bold=True, name="SimHei")

    rows = [
        ("实验1", "个人信息卡片 / 本地静态社交分享平台", ""),
        ("1.1", "任务目的", ""),
        ("1.2", "实验要求", ""),
        ("1.3", "实验过程与结果", ""),
        ("1.4", "心得体会", ""),
        ("实验2", "简易版微博 / ShareFlow 社交分享平台", ""),
        ("2.1", "任务目的", ""),
        ("2.2", "功能介绍", ""),
        ("2.3", "实验要求", ""),
        ("2.4", "实验过程与结果", ""),
        ("2.5", "心得体会", ""),
        ("附录", "AI辅助过程记录与指导教师评语", ""),
    ]
    for no, title, page in rows:
        cells = table.add_row().cells
        cells[0].text = no
        cells[1].text = title
        cells[2].text = page
        for i, cell in enumerate(cells):
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    set_run_font(run, size=10.5)

    note = doc.add_paragraph("说明：正式打印或转 PDF 前，可在 Word 中根据最终分页补充页码。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in note.runs:
        set_run_font(run, size=10, color=(100, 116, 139))
    doc.add_page_break()


def experiment_one(doc):
    add_heading(doc, "实验1：个人信息卡片 / 本地静态社交分享平台", 1)
    add_heading(doc, "一、任务目的", 2)
    add_text(doc, "本实验以浏览器可直接打开的本地网页为基础，综合运用 HTML、CSS 与 JavaScript 完成一个具有信息展示、内容发布、图片预览、评论、点赞和权限演示能力的前端应用。通过该实验，掌握语义化页面结构、响应式布局、组件化样式拆分、DOM 事件处理、本地存储和前端权限控制的基本方法，为后续使用 Vue 等现代前端框架开发完整项目奠定基础。")
    add_heading(doc, "二、实验要求", 2)
    add_bullets(doc, [
        "使用 HTML 构建顶部导航、搜索框、主题区域、发布表单、动态列表和侧边栏等页面结构。",
        "使用 CSS 实现社区内容流布局、按钮和表单状态、图片网格、响应式适配以及清晰的视觉层级。",
        "使用 JavaScript 完成动态发布、图片上传预览、按话题筛选、关键字搜索、评论、点赞、删除和身份切换。",
        "使用 localStorage 保存本地发布内容、评论与点赞状态，使刷新页面后数据仍然存在。",
        "区分游客、普通用户、管理员三种身份，体现不同操作权限。"
    ])
    add_heading(doc, "三、实验过程与结果", 2)
    add_heading(doc, "1. 项目结构设计", 3)
    add_text(doc, "实验一项目位于 C:\\Users\\。\\Desktop\\web，是一个无需安装依赖、无需启动服务端的本地静态网页。项目文件保持简洁：index.html 负责页面结构，styles.css 负责页面样式与响应式布局，app.js 负责交互逻辑、权限判断和本地数据读写，Agent.md 记录项目方案与后续扩展思路。")
    add_table(doc, ["文件", "作用"], [
        ["index.html", "定义 ShareHub 页面骨架、SVG 图标、发布表单、动态卡片模板和侧栏结构。"],
        ["styles.css", "实现三栏社区布局、按钮/表单样式、图片网格、移动端单列适配。"],
        ["app.js", "维护帖子数据、身份状态、搜索筛选、发布评论点赞删除等交互逻辑。"],
        ["localStorage", "使用 sharehub-local-posts 键保存帖子、评论、点赞和图片 Base64 数据。"],
    ])
    doc.add_picture(str(ASSET_DIR / "experiment1_structure.png"), width=Cm(15.8))
    add_heading(doc, "2. 页面结构实现", 3)
    add_text(doc, "index.html 使用 header、main、aside、section、template 等结构组织内容。顶部区域包含品牌、搜索框和身份切换控件；中间内容区包含介绍区域、发布表单和动态流；左侧提供热门话题与权限说明，右侧展示社区统计和后续扩展项。动态卡片使用 template 作为复用模板，由 JavaScript 克隆后填充具体数据。")
    add_code(doc, """
<form id="postForm">
  <textarea id="postText" maxlength="500"></textarea>
  <input id="imageInput" type="file" accept="image/png,image/jpeg,image/webp" multiple />
  <select id="visibilitySelect">
    <option value="public">公开</option>
    <option value="private">仅自己</option>
  </select>
</form>
""")
    add_heading(doc, "3. 交互逻辑实现", 3)
    add_text(doc, "app.js 中通过 posts、role、activeTopic、searchTerm、selectedImages 等变量维护页面状态。loadPosts() 从 localStorage 读取历史数据，savePosts() 在发布、删除、评论、点赞后同步保存。canWrite() 与 canDeletePost() 根据当前身份判断用户是否可以发布、评论、点赞或删除内容。")
    add_text(doc, "图片上传使用 FileReader 将本地图片转换为 Data URL，再在发布前放入 selectedImages 数组并渲染预览网格；发布成功后新动态被加入 posts 并重新渲染信息流。搜索和话题筛选通过 filteredPosts() 同时组合身份可见性、话题和关键词过滤，使游客只能看到公开内容。")
    add_heading(doc, "4. 实验结果", 3)
    add_text(doc, "实验一最终实现了一个可直接双击 index.html 打开的本地社交分享网页。页面支持浏览示例动态、按话题筛选、搜索作者/话题/内容、切换游客/普通用户/管理员身份、发布文字和图片动态、评论、点赞、删除以及刷新后保持本地数据。该实验虽然没有真实后端和数据库，但完整模拟了社区型 Web 应用的前端主要交互流程。")
    add_heading(doc, "四、心得体会", 2)
    add_text(doc, "通过实验一，我对前端三件套的分工有了更清晰的理解。HTML 负责结构，CSS 负责视觉层级与响应式布局，JavaScript 负责状态变化和用户交互。相比只做静态页面，本实验中发布、评论、点赞和权限控制都需要数据驱动页面重新渲染，这让我体会到前端状态管理的重要性。")
    add_text(doc, "实验中遇到的主要问题是如何在不使用后端的情况下模拟数据持久化和权限差异。最终使用 localStorage 保存动态数据，并通过身份选择器模拟游客、普通用户和管理员的操作差异。这个过程让我认识到：本地原型适合快速验证交互逻辑，但真正上线时仍需要将身份认证、权限校验、数据存储和文件服务迁移到服务端完成。")


def experiment_two(doc):
    add_heading(doc, "实验2：简易版微博 / ShareFlow 社交分享平台", 1)
    add_heading(doc, "一、任务目的", 2)
    add_text(doc, "实验二在实验一原型的基础上进一步完成前后端分离的简易版微博平台。项目位于 C:\\Users\\。\\Desktop\\webwork，前端采用 Vue3、Vite、Vue Router、Pinia 与 Axios，后端采用 Java 8、Spring Boot 2.7、Spring MVC、MyBatis 与 MySQL。实验目标是掌握现代 Web 应用的分层架构、REST API 设计、路由守卫、统一请求封装、数据库表设计、文件上传、通知消息和权限控制等完整开发流程。")
    add_heading(doc, "二、功能介绍", 2)
    add_text(doc, "ShareFlow 面向校园或兴趣社区中的图文分享场景，提供从用户登录、内容发布、互动反馈到后台管理的一组完整功能。平台首页展示公开信息流，用户登录后可以发布图文动态、进入详情页评论与点赞，也可以查看个人主页、接收通知、发送私信；管理员则可以进入后台查看统计数据并管理平台内容。")
    add_table(doc, ["功能模块", "功能说明", "涉及页面/接口"], [
        ["用户认证", "支持注册、登录、退出和当前用户信息获取，前端使用 shareflow_token 保持登录状态。", "LoginView、RegisterView、/api/auth/*"],
        ["信息流浏览", "首页按时间展示动态卡片，包含作者、标题、正文、图片、标签、点赞数和评论数。", "HomeView、PostCard、GET /api/posts"],
        ["内容发布", "登录用户可发布标题、正文、标签和图片，上传文件后由后端保存并返回访问地址。", "PublishView、PostForm、ImageUploader、POST /api/posts"],
        ["帖子详情", "展示单条动态的完整内容和评论区，用户可在详情页继续互动。", "PostDetailView、GET /api/posts/{id}"],
        ["评论与点赞", "支持发表评论、删除评论、点赞和取消点赞，后端通过唯一约束避免重复点赞。", "CommentList、LikeButton、/api/comments、/api/posts/{postId}/like"],
        ["个人主页", "展示用户资料和该用户发布的动态，支持当前用户编辑个人资料。", "ProfileView、GET /api/users/{id}"],
        ["通知与私信", "点赞、评论和私信事件可形成站内通知，用户可以查看未读摘要和私信会话。", "NotificationsView、MessagesView、/api/notifications、/api/messages"],
        ["后台管理", "管理员可查看统计数据和帖子列表，对异常内容进行管理。", "AdminView、/api/admin/stats、/api/admin/posts"],
    ])
    add_heading(doc, "三、实验要求", 2)
    add_bullets(doc, [
        "实现用户注册、登录、退出、当前用户信息获取和基于 token 的前端会话保持。",
        "实现内容信息流、发布动态、编辑删除动态、帖子详情、评论、点赞和图片上传。",
        "实现个人主页、通知中心、私信会话、管理员统计与内容管理等扩展功能。",
        "前端使用 Vue Router 管理页面路由，使用 Pinia 管理用户状态、信息流和消息通知状态。",
        "后端采用 Controller、Service、Mapper、Entity、DTO、VO 等分层组织代码，通过 MyBatis 访问 MySQL。",
        "数据库至少包含用户、帖子、评论、点赞、文件、私信、会话隐藏、通知等核心数据表。"
    ])
    add_heading(doc, "四、实验过程与结果", 2)
    add_heading(doc, "1. 总体架构设计", 3)
    add_text(doc, "ShareFlow 采用前后端分离结构。前端运行在 Vite 开发服务器（默认 http://localhost:5173），通过 /api 和 /uploads 代理访问后端；后端运行在 Spring Boot（默认 http://localhost:8080），负责统一接口、业务规则、权限校验、文件上传和数据库读写；MySQL 数据库 social_share 存储用户、帖子、评论、点赞、文件、私信和通知数据。")
    doc.add_picture(str(ASSET_DIR / "experiment2_flow.png"), width=Cm(15.8))
    add_table(doc, ["层次", "目录/文件", "职责"], [
        ["前端页面", "frontend/src/views", "HomeView、PublishView、PostDetailView、ProfileView、AdminView 等页面。"],
        ["前端组件", "frontend/src/components", "PostCard、PostForm、CommentList、ImageUploader、LikeButton、AppHeader 等复用组件。"],
        ["前端状态", "frontend/src/stores", "userStore、postStore、inboxStore 管理登录状态、帖子数据和未读消息。"],
        ["前端接口", "frontend/src/api", "Axios 实例与 authApi、postApi、commentApi、messageApi 等接口封装。"],
        ["后端控制器", "backend/src/main/java/.../controller", "Auth、Post、Comment、Like、File、User、Message、Notification、Admin 接口入口。"],
        ["后端业务", "backend/src/main/java/.../service", "处理认证、权限、发布、评论、点赞、文件、通知和私信逻辑。"],
        ["数据访问", "backend/src/main/resources/mapper", "MyBatis XML 映射 SQL，与 mapper 接口配合访问数据库。"],
    ])
    add_heading(doc, "2. 前端实现", 3)
    add_text(doc, "前端使用 Vue3 + Vite 构建 SPA。router/index.js 定义了首页、登录、注册、发布、帖子详情、个人主页、通知、私信和管理员页面。路由守卫在进入需要登录的页面前调用 userStore.fetchMe() 初始化登录状态；未登录访问发布、通知、私信等页面时跳转到登录页；非管理员访问 /admin 时返回首页。")
    add_code(doc, """
router.beforeEach(async (to) => {
  const userStore = useUserStore()
  if (!userStore.initialized) await userStore.fetchMe()
  if (to.meta.requiresAuth && !userStore.isLoggedIn) {
    return { name: 'login', query: { redirect: to.fullPath } }
  }
  if (to.meta.admin && !userStore.isAdmin) return { name: 'home' }
  return true
})
""")
    add_text(doc, "userStore 使用 localStorage 中的 shareflow_token 保存登录凭证，并在登录后刷新通知与私信摘要。api/http.js 创建 Axios 实例，统一设置 baseURL、请求超时、Authorization 请求头和响应体解析逻辑，使各业务 API 文件只关注具体接口路径。")
    add_heading(doc, "3. 后端接口实现", 3)
    add_text(doc, "后端基于 Spring Boot 2.7 与 Java 8 实现。控制器层提供 REST 风格接口：/api/auth 负责注册、登录、退出和当前用户；/api/posts 负责帖子列表、发布、详情、更新和删除；/api/posts/{postId}/comments 负责评论；/api/posts/{postId}/like 负责点赞和取消点赞；/api/files/upload 负责图片上传；/api/messages 与 /api/notifications 负责私信和通知；/api/admin 负责后台统计与内容管理。")
    add_table(doc, ["模块", "主要接口", "说明"], [
        ["认证", "POST /api/auth/register, POST /api/auth/login, GET /api/auth/me", "完成注册、登录和当前用户识别。"],
        ["帖子", "GET/POST /api/posts, GET/PUT/DELETE /api/posts/{id}", "完成信息流、发布、详情、编辑和删除。"],
        ["评论", "GET/POST /api/posts/{postId}/comments, DELETE /api/comments/{id}", "完成评论列表、发表评论和删除评论。"],
        ["点赞", "POST/DELETE /api/posts/{postId}/like", "完成点赞与取消点赞。"],
        ["文件", "POST /api/files/upload, GET /api/files/{id}", "上传图片并记录文件元数据。"],
        ["私信/通知", "GET/POST /api/messages, GET /api/notifications", "实现用户之间的私信会话和站内提醒。"],
        ["管理", "GET /api/admin/stats, GET /api/admin/posts", "提供管理员统计和内容管理能力。"],
    ])
    add_heading(doc, "4. 数据库设计", 3)
    add_text(doc, "数据库初始化脚本位于 backend/src/main/resources/schema.sql。系统使用 social_share 数据库，核心表包括 users、posts、comments、likes、files、direct_messages、message_thread_hides 和 notifications。表之间通过外键建立用户与帖子、帖子与评论、帖子与点赞、用户与文件、私信发送接收双方、通知关联内容之间的关系。")
    add_table(doc, ["数据表", "主要字段", "作用"], [
        ["users", "username、nickname、password_hash、avatar_url、bio、role、status", "存储用户基础资料、密码哈希、角色与状态。"],
        ["posts", "user_id、title、content、image_url、tags、visibility", "存储用户发布的动态内容、图片和可见范围。"],
        ["comments", "post_id、user_id、content", "存储帖子评论，并在帖子删除时级联删除。"],
        ["likes", "post_id、user_id", "通过唯一约束避免同一用户重复点赞同一帖子。"],
        ["files", "original_name、storage_path、url、content_type、size", "保存上传文件元信息和访问地址。"],
        ["direct_messages", "sender_id、receiver_id、content、read_at", "存储用户私信内容与阅读状态。"],
        ["notifications", "user_id、actor_id、type、post_id、comment_id、message_id、is_read", "存储点赞、评论、私信等事件提醒。"],
    ])
    add_heading(doc, "5. 测试与运行结果", 3)
    add_text(doc, "项目提供前后端启动脚本：scripts/start-frontend.ps1 启动 Vite 前端，scripts/start-backend.ps1 启动 Spring Boot 后端，scripts/init-db.ps1 可初始化 MySQL 数据库。后端 target/surefire-reports 中保留了 AuthServiceTest、PostServiceTest、LikeServiceTest、FileServiceTest、MessageServiceTest、Notification 相关测试等报告文件，说明核心业务逻辑已经纳入单元测试范围。")
    add_text(doc, "从功能结果看，实验二已经从实验一的本地静态原型升级为完整前后端分离项目：用户可注册登录，登录后发布图文动态、浏览信息流、查看帖子详情、评论点赞、访问个人主页、接收通知和私信；管理员可进入后台页面查看统计和管理内容。前端负责交互体验，后端负责真实数据存储和最终权限校验。")
    add_heading(doc, "五、心得体会", 2)
    add_text(doc, "实验二让我体会到完整 Web 项目与静态页面原型之间的差异。前端不再只是操作 DOM，而是需要路由、状态、请求封装、错误处理和权限提示共同协作；后端也不只是返回数据，而要处理参数校验、业务规则、数据库一致性、文件存储和异常响应。")
    add_text(doc, "在开发过程中，路由守卫和 token 会话保持是一个重点。前端必须在进入受保护页面前确认当前用户状态，请求接口时自动携带 Authorization 信息；一旦 token 失效，需要清理本地状态并引导用户重新登录。后端方面，点赞唯一约束、评论级联删除、通知未读统计、私信会话隐藏等逻辑都体现了数据库设计与业务规则之间的紧密联系。")
    add_text(doc, "通过本实验，我对前后端分离架构、REST 接口、组件化页面、状态管理、MyBatis 数据访问和 MySQL 表关系有了系统认识。实验一让我完成了交互原型，实验二则将原型转化为可持久化、可扩展、具备权限控制的课程项目，这一过程提升了我的工程化思维和综合调试能力。")


def appendix(doc):
    add_heading(doc, "附录：AI辅助过程记录与指导教师评语", 1)
    add_text(doc, "本报告撰写过程中使用 AI 辅助梳理项目结构、归纳源码证据和生成结构图。报告内容仍以本机项目文件为依据，包括 C:\\Users\\。\\Desktop\\web 的 index.html、styles.css、app.js，以及 C:\\Users\\。\\Desktop\\webwork 的 frontend/src、backend/src、schema.sql、pom.xml、package.json 和 surefire-reports 等文件。")
    doc.add_picture(str(ASSET_DIR / "ai_prompt_evidence.png"), width=Cm(15.8))
    add_text(doc, "指导教师评语：")
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = "\n\n\n\n\n\n       签名：                         年    月    日"


def validate_docx(path):
    required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
    with ZipFile(path) as zf:
        names = set(zf.namelist())
        missing = required - names
        media = [n for n in names if n.startswith("word/media/")]
    if missing:
        raise RuntimeError(f"missing docx parts: {missing}")
    if len(media) < 3:
        raise RuntimeError(f"expected embedded media, got {len(media)}")


def main():
    build_assets()
    doc = Document()
    style_document(doc)
    cover(doc)
    toc(doc)
    experiment_one(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    experiment_two(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    appendix(doc)
    doc.save(OUT)
    validate_docx(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
